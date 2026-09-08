"""Comprehensive test suite for Phase 10 - Question Paper Intelligence.

Tests document extraction (PDF, DOCX, TXT), scanned PDF detection,
question parsing, marks extraction, NLP topic clustering, rule classification,
difficulty estimation, and Flask role-based access control.
"""

import os
import unittest
from pathlib import Path
import tempfile
import pymupdf as fitz
import docx

from app import app
from question_paper.extractor import extract_text_from_file, extract_from_pdf, extract_from_docx, extract_from_txt
from question_paper.parser import parse_question_paper, extract_marks_from_text, extract_declared_total_marks
from question_paper.analyzer import (
    classify_question_type,
    estimate_question_difficulty,
    extract_keywords_and_clusters,
    analyze_question_paper
)


class TestQuestionPaperIntelligence(unittest.TestCase):
    
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.mkdtemp()
        
        # 1. Create a Test TXT paper
        cls.txt_path = Path(cls.temp_dir) / "test_paper.txt"
        with open(cls.txt_path, "w", encoding="utf-8") as f:
            f.write(
                "UNIVERSITY EXAMINATION\n"
                "Max Marks: 50\n\n"
                "Q1. (a) Define Machine Learning and state its types. [5]\n"
                "Q1. (b) Explain Linear Regression with cost function. [5]\n"
                "Q2. (a) Write Python code to implement Decision Tree using scikit-learn. [10]\n"
                "Q2. (b) Calculate the Euclidean distance between points (2, 3) and (5, 7). [5]\n"
                "Q3. (a) Design an end-to-end spam classification pipeline. [10]\n"
                "Q3. (b) Differentiate between Supervised and Unsupervised Learning. [5]\n"
            )

        # 2. Create a Test DOCX paper
        cls.docx_path = Path(cls.temp_dir) / "test_paper.docx"
        doc = docx.Document()
        doc.add_heading("UNIVERSITY SEMESTER EXAM", level=1)
        doc.add_paragraph("Total Marks: 30")
        doc.add_paragraph("1. Define Normalization and explain 1NF. (5 Marks)")
        doc.add_paragraph("2. Write SQL query to join Employee and Department tables. (10 Marks)")
        doc.add_paragraph("3. Compute the time complexity of Quick Sort in worst case. [5]")
        doc.add_paragraph("4. Design an ER diagram for University Library System. [10]")
        doc.save(str(cls.docx_path))

        # 3. Create a Test PDF paper
        cls.pdf_path = Path(cls.temp_dir) / "test_paper.pdf"
        pdf_doc = fitz.open()
        page = pdf_doc.new_page()
        page.insert_text(
            (50, 72),
            "INSTITUTE OF TECHNOLOGY\n"
            "Maximum Marks: 40\n\n"
            "Q1. What is an Operating System? List its primary functions. [5]\n"
            "Q2. Explain CPU Scheduling algorithms: FCFS and Round Robin. [10]\n"
            "Q3. Write a program to create a child process using fork() system call. [10]\n"
            "Q4. Calculate the average waiting time for given processes with arrival times. [5]\n"
            "Q5. Explain Deadlock conditions and Banker's algorithm in detail. [10]\n"
        )
        pdf_doc.save(str(cls.pdf_path))
        pdf_doc.close()

        # 4. Create a Scanned/Empty PDF paper (simulating image-only scan with no text)
        cls.scanned_pdf_path = Path(cls.temp_dir) / "scanned_paper.pdf"
        scanned_doc = fitz.open()
        scanned_doc.new_page()  # Blank page without text
        scanned_doc.save(str(cls.scanned_pdf_path))
        scanned_doc.close()

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(cls.temp_dir, ignore_errors=True)

    # -------------------------------------------------------------
    # 1. Extraction Tests
    # -------------------------------------------------------------
    def test_txt_extraction(self):
        result = extract_text_from_file(self.txt_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "TXT")
        self.assertIn("Machine Learning", result["text"])

    def test_docx_extraction(self):
        result = extract_text_from_file(self.docx_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "DOCX")
        self.assertIn("Normalization", result["text"])

    def test_pdf_extraction(self):
        result = extract_text_from_file(self.pdf_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "PDF")
        self.assertIn("Operating System", result["text"])

    def test_scanned_pdf_detection(self):
        result = extract_text_from_file(self.scanned_pdf_path)
        self.assertFalse(result["success"])
        self.assertTrue(result["is_scanned"])
        self.assertIn("scanned/image-based", result["error"])

    def test_unsupported_extension(self):
        fake_path = Path(self.temp_dir) / "test.xyz"
        with open(fake_path, "w") as f:
            f.write("test")
        result = extract_text_from_file(fake_path)
        self.assertFalse(result["success"])
        self.assertIn("Unsupported file extension", result["error"])

    # -------------------------------------------------------------
    # 2. Parsing & Marks Detection Tests
    # -------------------------------------------------------------
    def test_marks_extraction(self):
        m1, t1 = extract_marks_from_text("Define DBMS and its features. [5]")
        self.assertEqual(m1, 5)
        self.assertEqual(t1, "Define DBMS and its features.")

        m2, t2 = extract_marks_from_text("Explain Normalization in detail. (10 Marks)")
        self.assertEqual(m2, 10)
        self.assertEqual(t2, "Explain Normalization in detail.")

        m3, t3 = extract_marks_from_text("What is a Database?")
        self.assertIsNone(m3)
        self.assertEqual(t3, "What is a Database?")

    def test_declared_total_marks(self):
        text = "UNIVERSITY EXAM\nSubject: DBMS\nMax Marks: 70\nTime: 3 Hours"
        self.assertEqual(extract_declared_total_marks(text), 70)

        text2 = "College Exam\nTotal Marks: 100\nDate: 2026"
        self.assertEqual(extract_declared_total_marks(text2), 100)

    def test_full_question_parsing(self):
        txt_res = extract_text_from_file(self.txt_path)
        parsed = parse_question_paper(txt_res["text"])
        self.assertTrue(parsed["success"])
        self.assertEqual(parsed["total_questions"], 6)
        self.assertEqual(parsed["declared_total_marks"], 50)
        self.assertEqual(parsed["total_detected_marks"], 40)  # 5+5+10+5+10+5 = 40
        self.assertEqual(parsed["questions_with_marks"], 6)

    # -------------------------------------------------------------
    # 3. NLP Analysis, Classification & Difficulty Tests
    # -------------------------------------------------------------
    def test_question_type_classification(self):
        self.assertEqual(classify_question_type("Define Polymorphism and state its types.", 5), "Definition")
        self.assertEqual(classify_question_type("Write Python code to implement Binary Search.", 10), "Programming")
        self.assertEqual(classify_question_type("Calculate the determinant of matrix A.", 5), "Numerical")
        self.assertEqual(classify_question_type("Design an ER schema for an e-commerce platform.", 10), "Application")
        self.assertEqual(classify_question_type("Explain the working of TCP three-way handshake.", 5), "Theory/Descriptive")

    def test_difficulty_estimation(self):
        easy_diff = estimate_question_difficulty("Define DBMS and list 3 advantages.", 2, "Definition")
        self.assertEqual(easy_diff["difficulty"], "Easy")

        hard_diff = estimate_question_difficulty("Design and optimize an end-to-end distributed system architecture with fault tolerance.", 15, "Application")
        self.assertEqual(hard_diff["difficulty"], "Hard")

    def test_full_analysis_pipeline(self):
        txt_res = extract_text_from_file(self.txt_path)
        parsed = parse_question_paper(txt_res["text"])
        analysis = analyze_question_paper(parsed)

        self.assertTrue(analysis["success"])
        self.assertIn("kpis", analysis)
        self.assertIn("charts", analysis)
        self.assertIn("topic_clusters", analysis)
        self.assertIn("insights", analysis)
        self.assertGreater(len(analysis["insights"]), 0)
        self.assertGreater(len(analysis["top_keywords"]), 0)

    # -------------------------------------------------------------
    # 4. Flask Security & Role-Based Authorization Tests
    # -------------------------------------------------------------
    def test_security_access_control(self):
        client = app.test_client()

        # A. Unauthenticated user -> redirected to login
        res = client.get("/question-paper", follow_redirects=False)
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login", res.headers["Location"])

        # B. Student user -> blocked and redirected to student-dashboard
        with client.session_transaction() as sess:
            sess["student_logged_in"] = True
            sess["student_roll"] = "101"
            sess["student_name"] = "Alice"

        res_student = client.get("/question-paper", follow_redirects=False)
        self.assertEqual(res_student.status_code, 302)
        self.assertIn("/student-dashboard", res_student.headers["Location"])

        # C. Admin user -> authorized to view portal
        with client.session_transaction() as sess:
            sess.clear()
            sess["admin_logged_in"] = True
            sess["admin_id"] = "T001"
            sess["admin_name"] = "Mr. Zeeshan Shaikh"

        res_admin = client.get("/question-paper")
        self.assertEqual(res_admin.status_code, 200)
        self.assertIn(b"Question Paper Intelligence", res_admin.data)

        # D. Teacher user -> authorized to view assigned courses
        with client.session_transaction() as sess:
            sess.clear()
            sess["teacher_logged_in"] = True
            sess["teacher_username"] = "prof_amit"
            sess["teacher_name"] = "Prof. Amit Sharma"
            sess["teacher_id"] = "T002"

        res_teacher = client.get("/question-paper")
        self.assertEqual(res_teacher.status_code, 200)

        # E. Teacher unauthorized subject tampering test:
        # prof_amit is assigned Data Structures (Sem 2) and Python (Sem 1).
        # Attempting to analyze "Machine Learning" (Sem 5) must be blocked.
        with open(self.txt_path, "rb") as f:
            res_tamper = client.post(
                "/question-paper/analyze",
                data={
                    "semester": "5",
                    "subject": "Machine Learning",
                    "paper_file": (f, "test_paper.txt")
                },
                content_type="multipart/form-data"
            )
        self.assertEqual(res_tamper.status_code, 200)
        self.assertIn(b"Authorization Error", res_tamper.data)


if __name__ == "__main__":
    unittest.main()
