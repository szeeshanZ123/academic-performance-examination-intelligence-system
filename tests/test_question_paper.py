"""Comprehensive test suite for Phase 8.3 - Question Paper Intelligence Finalization.

Tests document extraction (PDF, DOCX, TXT), scanned PDF detection,
instruction skipping, question and subquestion parsing, marks extraction,
total marks validation, syllabus topic mapping, Bloom's Taxonomy,
question type classification, difficulty estimation, quality checks,
and Flask role-based access control (teachers restricted to assignments, students blocked).
"""

import os
import unittest
from pathlib import Path
import tempfile
import pymupdf as fitz
import docx

from app import app
from question_paper.extractor import extract_text_from_file, extract_from_pdf, extract_from_docx, extract_from_txt
from question_paper.parser import parse_question_paper, extract_marks_from_text, extract_declared_total_marks, is_instruction_line
from question_paper.analyzer import (
    classify_question_type,
    classify_bloom_level,
    classify_topic,
    estimate_question_difficulty,
    extract_keywords_and_clusters,
    analyze_question_paper
)


class TestQuestionPaperIntelligenceFinalization(unittest.TestCase):
    
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.mkdtemp()
        
        # 1. Create a Test TXT paper with Instructions and Subquestions
        cls.txt_path = Path(cls.temp_dir) / "test_paper.txt"
        with open(cls.txt_path, "w", encoding="utf-8") as f:
            f.write(
                "UNIVERSITY EXAMINATION\n"
                "Max Marks: 50\n"
                "Instructions:\n"
                "1. Answer all questions.\n"
                "2. Figures to the right indicate full marks.\n"
                "3. Assume suitable data wherever necessary.\n\n"
                "Q1. (a) Define Machine Learning and state its types. [5]\n"
                "Q1. (b) Explain Linear Regression with cost function. [5]\n"
                "Q2. (a) Write Python code to implement Decision Tree using scikit-learn. [10]\n"
                "Q2. (b) Calculate the Euclidean distance between points (2, 3) and (5, 7). [5]\n"
                "Q3. (a) Design an end-to-end spam classification pipeline. [10]\n"
                "Q3. (b) Differentiate between Supervised and Unsupervised Learning. [5]\n"
            )

        # 2. Create a Test DOCX paper
        cls.docx_path = Path(cls.temp_dir) / "test_paper.docx"
        doc = docx.Document()
        doc.add_heading("UNIVERSITY SEMESTER EXAM", level=1)
        doc.add_paragraph("Total Marks: 30")
        doc.add_paragraph("1. Define Normalization and explain 1NF. (5 Marks)")
        doc.add_paragraph("2. Write SQL query to join Employee and Department tables. (10 Marks)")
        doc.add_paragraph("3. Compute the time complexity of Quick Sort in worst case. [5]")
        doc.add_paragraph("4. Design an ER diagram for University Library System. [10]")
        doc.save(str(cls.docx_path))

        # 3. Create a Test PDF paper
        cls.pdf_path = Path(cls.temp_dir) / "test_paper.pdf"
        pdf_doc = fitz.open()
        page = pdf_doc.new_page()
        page.insert_text(
            (50, 72),
            "INSTITUTE OF TECHNOLOGY\n"
            "Maximum Marks: 40\n\n"
            "Q1. What is an Operating System? List its primary functions. [5]\n"
            "Q2. Explain CPU Scheduling algorithms: FCFS and Round Robin. [10]\n"
            "Q3. Write a program to create a child process using fork() system call. [10]\n"
            "Q4. Calculate the average waiting time for given processes with arrival times. [5]\n"
            "Q5. Explain Deadlock conditions and Banker's algorithm in detail. [10]\n"
        )
        pdf_doc.save(str(cls.pdf_path))
        pdf_doc.close()

        # 4. Create a Scanned/Empty PDF paper (simulating image-only scan with no text)
        cls.scanned_pdf_path = Path(cls.temp_dir) / "scanned_paper.pdf"
        scanned_doc = fitz.open()
        scanned_doc.new_page()  # Blank page without text
        scanned_doc.save(str(cls.scanned_pdf_path))
        scanned_doc.close()

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(cls.temp_dir, ignore_errors=True)

    # -------------------------------------------------------------
    # 1. Extraction Tests (PDF, DOCX, TXT, Scanned)
    # -------------------------------------------------------------
    def test_txt_extraction(self):
        result = extract_text_from_file(self.txt_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "TXT")
        self.assertIn("Machine Learning", result["text"])

    def test_docx_extraction(self):
        result = extract_text_from_file(self.docx_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "DOCX")
        self.assertIn("Normalization", result["text"])

    def test_pdf_extraction(self):
        result = extract_text_from_file(self.pdf_path)
        self.assertTrue(result["success"])
        self.assertEqual(result["file_type"], "PDF")
        self.assertIn("Operating System", result["text"])

    def test_scanned_pdf_detection(self):
        result = extract_text_from_file(self.scanned_pdf_path)
        self.assertFalse(result["success"])
        self.assertTrue(result["is_scanned"])
        self.assertIn("scanned/image-based", result["error"])

    def test_unsupported_extension(self):
        fake_path = Path(self.temp_dir) / "test.xyz"
        with open(fake_path, "w") as f:
            f.write("test")
        result = extract_text_from_file(fake_path)
        self.assertFalse(result["success"])
        self.assertIn("Unsupported file extension", result["error"])

    # -------------------------------------------------------------
    # 2. Instruction Skipping & Marks Parsing Tests
    # -------------------------------------------------------------
    def test_instruction_line_detection(self):
        self.assertTrue(is_instruction_line("1. Answer all questions."))
        self.assertTrue(is_instruction_line("Figures to the right indicate full marks."))
        self.assertTrue(is_instruction_line("Assume suitable data wherever necessary."))
        self.assertTrue(is_instruction_line("Time: 3 Hours"))
        self.assertTrue(is_instruction_line("Max Marks: 70"))
        self.assertFalse(is_instruction_line("Explain Pandas DataFrame and Series. [5]"))

    def test_instruction_skipping_in_parser(self):
        txt_res = extract_text_from_file(self.txt_path)
        parsed = parse_question_paper(txt_res["text"])
        self.assertTrue(parsed["success"])
        # Exactly 6 academic questions should be parsed; instructions 1., 2., 3. must NOT be questions
        self.assertEqual(parsed["total_questions"], 6)
        question_numbers = [q["question_number"] for q in parsed["questions"]]
        self.assertEqual(question_numbers, ["Q1(a)", "Q1(b)", "Q2(a)", "Q2(b)", "Q3(a)", "Q3(b)"])

    def test_marks_extraction_patterns(self):
        m1, t1 = extract_marks_from_text("Define DBMS and its features. [5]")
        self.assertEqual(m1, 5)
        self.assertEqual(t1, "Define DBMS and its features.")

        m2, t2 = extract_marks_from_text("Explain Normalization in detail. (10 Marks)")
        self.assertEqual(m2, 10)
        self.assertEqual(t2, "Explain Normalization in detail.")

        m3, t3 = extract_marks_from_text("What is a Database?")
        self.assertIsNone(m3)
        self.assertEqual(t3, "What is a Database?")

        m4, t4 = extract_marks_from_text("Write a Python script to filter data. Marks: 10")
        self.assertEqual(m4, 10)

    def test_declared_total_marks_detection(self):
        text = "UNIVERSITY EXAM\nSubject: DBMS\nMax Marks: 70\nTime: 3 Hours"
        self.assertEqual(extract_declared_total_marks(text), 70)

        text2 = "College Exam\nTotal Marks: 100\nDate: 2026"
        self.assertEqual(extract_declared_total_marks(text2), 100)

    # -------------------------------------------------------------
    # 3. Topic Mapping & "Other / Unclassified" Fallback Tests
    # -------------------------------------------------------------
    def test_topic_classification(self):
        # Python for Data Analytics topics
        t1 = classify_topic("Explain NumPy ndarray, indexing, and slicing.", "Python for Data Analytics")
        self.assertEqual(t1, "NumPy")

        t2 = classify_topic("Write a Python script using Pandas dataframe and groupby().", "Python for Data Analytics")
        self.assertEqual(t2, "Pandas")

        t3 = classify_topic("Handle missing values and feature scaling using scikit-learn.", "Python for Data Analytics")
        self.assertEqual(t3, "Data Cleaning")

        t4 = classify_topic("Plot boxplots and correlation heatmap using Seaborn.", "Python for Data Analytics")
        self.assertEqual(t4, "Data Visualization")

        # Unrelated question -> Must map to "Other / Unclassified"
        t_unrelated = classify_topic("Explain French Revolution and European history.", "Python for Data Analytics")
        self.assertEqual(t_unrelated, "Other / Unclassified")

    # -------------------------------------------------------------
    # 4. Bloom's Taxonomy & Question Type Tests
    # -------------------------------------------------------------
    def test_bloom_taxonomy_classification(self):
        self.assertEqual(classify_bloom_level("Define Polymorphism and state its types."), "Remember")
        self.assertEqual(classify_bloom_level("Explain the working of TCP three-way handshake."), "Understand")
        self.assertEqual(classify_bloom_level("Write a Python script to calculate correlation."), "Apply")
        self.assertEqual(classify_bloom_level("Differentiate between EDA and Confirmatory Data Analysis."), "Analyze")
        self.assertEqual(classify_bloom_level("Evaluate and justify the choice of BCNF over 3NF."), "Evaluate")
        self.assertEqual(classify_bloom_level("Design an end-to-end data visualization pipeline."), "Create")

    def test_question_type_classification(self):
        self.assertEqual(classify_question_type("Define Polymorphism and state its types.", 5), "Definition")
        self.assertEqual(classify_question_type("Write Python code to implement Binary Search.", 10), "Programming")
        self.assertEqual(classify_question_type("Calculate the determinant of matrix A.", 5), "Numerical")
        self.assertEqual(classify_question_type("Design an ER schema for an e-commerce platform.", 10), "Application")
        self.assertEqual(classify_question_type("Compare and contrast supervised vs unsupervised learning.", 5), "Comparison")
        self.assertEqual(classify_question_type("Explain the working of TCP three-way handshake.", 5), "Explanation")

    def test_difficulty_estimation(self):
        easy_diff = estimate_question_difficulty("Define DBMS and list 3 advantages.", 2, "Definition", "Remember")
        self.assertEqual(easy_diff["difficulty"], "Easy")

        hard_diff = estimate_question_difficulty("Design and optimize an end-to-end distributed system architecture with fault tolerance.", 15, "Application", "Create")
        self.assertEqual(hard_diff["difficulty"], "Hard")

    # -------------------------------------------------------------
    # 5. Full Analysis Pipeline & Quality Checks Tests
    # -------------------------------------------------------------
    def test_full_analysis_pipeline_and_quality_checks(self):
        txt_res = extract_text_from_file(self.txt_path)
        parsed = parse_question_paper(txt_res["text"])
        analysis = analyze_question_paper(parsed, subject="Machine Learning")

        self.assertTrue(analysis["success"])
        self.assertIn("kpis", analysis)
        self.assertIn("charts", analysis)
        self.assertIn("topic_distribution", analysis)
        self.assertIn("quality_checks", analysis)
        self.assertIn("insights", analysis)
        self.assertGreater(len(analysis["insights"]), 0)
        self.assertGreater(len(analysis["quality_checks"]), 0)

        # Total marks check: 50 declared vs 40 detected -> Review Required
        self.assertEqual(analysis["kpis"]["declared_total_marks"], 50)
        self.assertEqual(analysis["kpis"]["total_detected_marks"], 40)
        self.assertEqual(analysis["kpis"]["marks_status"], "Review Required")

        # Verify each question has Bloom level and Topic
        for q in analysis["questions"]:
            self.assertIn("bloom_level", q)
            self.assertIn("topic", q)
            self.assertIn("question_type", q)
            self.assertIn("difficulty", q)

    # -------------------------------------------------------------
    # 6. Sample Paper End-to-End Tests (Semester 3 PDF)
    # -------------------------------------------------------------
    def test_sem3_demo_pdf_pipeline(self):
        sample_pdf = Path(__file__).resolve().parents[1] / "question_paper" / "samples" / "sem3_python_data_analytics.pdf"
        self.assertTrue(sample_pdf.exists())

        ext = extract_text_from_file(sample_pdf)
        self.assertTrue(ext["success"])

        parsed = parse_question_paper(ext["text"])
        self.assertTrue(parsed["success"])
        self.assertEqual(parsed["total_questions"], 11)

        analyzed = analyze_question_paper(parsed, subject="Python for Data Analytics")
        self.assertTrue(analyzed["success"])
        self.assertEqual(analyzed["kpis"]["total_questions"], 11)
        self.assertEqual(analyzed["kpis"]["total_detected_marks"], 75)
        self.assertEqual(analyzed["kpis"]["declared_total_marks"], 70)
        self.assertEqual(analyzed["kpis"]["marks_status"], "Review Required")

        # Verify topic distribution contains Pandas, NumPy, Data Cleaning, etc.
        topics = [t["topic"] for t in analyzed["topic_distribution"]]
        self.assertIn("Pandas", topics)
        self.assertIn("NumPy", topics)
        self.assertIn("Data Cleaning", topics)

    # -------------------------------------------------------------
    # 7. Flask Security & Role-Based Authorization Tests
    # -------------------------------------------------------------
    def test_security_access_control(self):
        client = app.test_client()

        # A. Unauthenticated user -> redirected to login
        res = client.get("/question-paper", follow_redirects=False)
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login", res.headers["Location"])

        # B. Student user -> blocked and redirected to student-dashboard
        with client.session_transaction() as sess:
            sess["student_logged_in"] = True
            sess["student_roll"] = "101"
            sess["student_name"] = "Alice"

        res_student = client.get("/question-paper", follow_redirects=False)
        self.assertEqual(res_student.status_code, 302)
        self.assertIn("/student-dashboard", res_student.headers["Location"])

        # Student also blocked from POST /question-paper/analyze
        res_student_post = client.post("/question-paper/analyze", follow_redirects=False)
        self.assertEqual(res_student_post.status_code, 302)
        self.assertIn("/student-dashboard", res_student_post.headers["Location"])

        # C. Admin user -> authorized to view portal
        with client.session_transaction() as sess:
            sess.clear()
            sess["admin_logged_in"] = True
            sess["admin_id"] = "T001"
            sess["admin_name"] = "Mr. Zeeshan Shaikh"

        res_admin = client.get("/question-paper")
        self.assertEqual(res_admin.status_code, 200)
        self.assertIn(b"Question Paper Intelligence", res_admin.data)

        # D. Teacher user -> authorized to view assigned courses
        with client.session_transaction() as sess:
            sess.clear()
            sess["teacher_logged_in"] = True
            sess["teacher_username"] = "prof_amit"
            sess["teacher_name"] = "Prof. Amit Sharma"
            sess["teacher_id"] = "T002"

        res_teacher = client.get("/question-paper")
        self.assertEqual(res_teacher.status_code, 200)

        # E. Teacher unauthorized subject tampering test:
        # prof_amit is assigned Data Structures (Sem 2) and Python Programming (Sem 1).
        # Attempting to analyze "Machine Learning" (Sem 5) must be blocked.
        with open(self.txt_path, "rb") as f:
            res_tamper = client.post(
                "/question-paper/analyze",
                data={
                    "semester": "5",
                    "subject": "Machine Learning",
                    "paper_file": (f, "test_paper.txt")
                },
                content_type="multipart/form-data"
            )
        self.assertEqual(res_tamper.status_code, 200)
        self.assertIn(b"Authorization Error", res_tamper.data)


if __name__ == "__main__":
    unittest.main()
